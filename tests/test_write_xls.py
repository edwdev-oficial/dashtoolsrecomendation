import unittest

import pandas as pd
from openpyxl import load_workbook

from dashtoolsrecomendation.services.write_xls import (
    gerar_excel,
    normalizar_coluna_idade_atual,
    normalizar_coluna_numero_serie,
)


class WriteXlsTests(unittest.TestCase):
    def test_consolida_aliases_do_numero_de_serie(self):
        source = pd.DataFrame(
            {
                "Id": ["1", "2"],
                "Número de série": ["", "123"],
                "Número de Série": ["456", "999"],
            }
        )

        result = normalizar_coluna_numero_serie(source)

        self.assertEqual(result.columns.tolist().count("Número de série"), 1)
        self.assertNotIn("Número de Série", result.columns)
        self.assertEqual(result["Número de série"].tolist(), ["456", "123"])

    def test_gera_excel_quando_as_duas_variacoes_estao_presentes(self):
        source = pd.DataFrame(
            {
                "Id": ["1"],
                "Número de série": [pd.NA],
                "Número de Série": ["456"],
            }
        )

        output = gerar_excel(source)
        exported = pd.read_excel(output)

        self.assertGreater(len(output.getvalue()), 0)
        self.assertEqual(exported.columns.tolist().count("Número de série"), 1)
        self.assertEqual(exported.loc[0, "Número de série"], 456)

    def test_padroniza_nome_da_coluna_idade_atual_no_relatorio(self):
        source = pd.DataFrame({"Idade atual": [4.5]})

        result = normalizar_coluna_idade_atual(source)
        output = gerar_excel(source)
        exported = pd.read_excel(output)

        self.assertNotIn("Idade atual", result.columns)
        self.assertIn("Idade Atual (anos)", result.columns)
        self.assertIn("Idade Atual (anos)", exported.columns)
        self.assertEqual(exported.loc[0, "Idade Atual (anos)"], 4.5)

    def test_centraliza_colunas_e_formata_idade_com_uma_casa_decimal(self):
        centered_columns = [
            "Idade atual",
            "Reparações no período",
            "Anos com reparação",
            "Garantia",
            "Fim período sem custo (dias)",
            "Índice de prioridade",
        ]
        source = pd.DataFrame(
            {
                "Descrição": ["Máquina de teste"],
                "Idade atual": [4.56],
                "Reparações no período": [3],
                "Anos com reparação": [2],
                "Garantia": ["Sim"],
                "Fim período sem custo (dias)": [45],
                "Índice de prioridade": [82.3],
                "Custo com impostos": [1234.5],
            }
        )

        workbook = load_workbook(gerar_excel(source))
        worksheet = workbook["Relatório"]
        headers = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        expected_headers = [
            "Idade Atual (anos)",
            *centered_columns[1:],
        ]
        for header in expected_headers:
            column = headers[header]
            self.assertEqual(worksheet.cell(1, column).alignment.horizontal, "center")
            self.assertEqual(worksheet.cell(2, column).alignment.horizontal, "center")

        age_column = headers["Idade Atual (anos)"]
        self.assertEqual(worksheet.cell(2, age_column).number_format, "0.0")

        cost_column = headers["Custo com impostos"]
        self.assertEqual(worksheet.cell(2, cost_column).number_format, "0.00")


if __name__ == "__main__":
    unittest.main()
